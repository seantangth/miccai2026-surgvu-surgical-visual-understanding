"""把 APC_TDC_presentation.pptx 與每張投影片的旁白音檔合成 1080p 影片。

流程：pptx 拆成單張 → qlmanage 渲染成 PNG（超取樣後降尺寸，字比較銳利）
      → 每張用「靜態圖 + 該張音檔」做一段 → 串接成一支 MP4。

用法：
    python make_video.py <deck.pptx> <audio_dir> <out.mp4> [--tail 0.4] [--super 3840]

<audio_dir> 內要有 slide1.* ～ slide4.*（mp3 / wav / m4a / aac 皆可），
檔名數字對應投影片順序。缺某一張就會直接報錯，不會靜默略過。
"""
import argparse
import json
import shutil
import struct
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation

W, H = 1920, 1080
AUDIO_EXT = (".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg")


def run(cmd, **kw):
    r = subprocess.run(cmd, capture_output=True, text=True, **kw)
    if r.returncode != 0:
        sys.exit(f"FAILED: {' '.join(map(str, cmd))}\n{r.stderr[-1500:]}")
    return r


def png_size(p):
    d = Path(p).read_bytes()[16:24]
    return struct.unpack(">II", d)


def duration(p):
    r = run(["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "json", str(p)])
    return float(json.loads(r.stdout)["format"]["duration"])


def render_slides(deck, workdir, supersample):
    """每張投影片各存成單張 pptx 再用 Quick Look 渲染（qlmanage 對多頁只出第一張）。"""
    n = len(Presentation(deck).slides)
    pngs = []
    for i in range(1, n + 1):
        prs = Presentation(deck)
        lst = prs.slides._sldIdLst
        for k, sid in enumerate(list(lst), 1):
            if k != i:
                lst.remove(sid)
        one = workdir / f"slide{i}.pptx"
        prs.save(one)
        run(["qlmanage", "-t", "-s", str(supersample), "-o", str(workdir), str(one)])
        raw = workdir / f"slide{i}.pptx.png"
        if not raw.exists():
            sys.exit(f"Quick Look 沒有產出 {raw}；請確認能開啟 pptx 預覽")
        out = workdir / f"frame{i}.png"
        # 超取樣後降到 1080p，長寬比不符時補白邊而不是裁切
        run(["ffmpeg", "-y", "-v", "error", "-i", str(raw),
             "-vf", f"scale={W}:{H}:force_original_aspect_ratio=decrease:flags=lanczos,"
                    f"pad={W}:{H}:(ow-iw)/2:(oh-ih)/2:white",
             str(out)])
        pngs.append(out)
        print(f"  slide {i}: {png_size(raw)[0]}x{png_size(raw)[1]} → {W}x{H}")
    return pngs


def find_audio(audio_dir, i):
    for ext in AUDIO_EXT:
        for name in (f"slide{i}{ext}", f"slide_{i}{ext}", f"{i}{ext}"):
            p = Path(audio_dir) / name
            if p.exists():
                return p
    sys.exit(f"找不到第 {i} 張的音檔：{audio_dir}/slide{i}.[{'/'.join(e[1:] for e in AUDIO_EXT)}]")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck"); ap.add_argument("audio_dir"); ap.add_argument("out")
    ap.add_argument("--tail", type=float, default=0.4, help="每張講完後多停留幾秒")
    ap.add_argument("--super", type=int, default=3840, dest="supersample")
    a = ap.parse_args()

    for tool in ("ffmpeg", "ffprobe", "qlmanage"):
        if not shutil.which(tool):
            sys.exit(f"缺少 {tool}")

    work = Path(tempfile.mkdtemp(prefix="surgvu_video_"))
    try:
        print("渲染投影片…")
        frames = render_slides(a.deck, work, a.supersample)

        print("合成分段…")
        segments, total = [], 0.0
        for i, frame in enumerate(frames, 1):
            au = find_audio(a.audio_dir, i)
            d = duration(au) + a.tail
            total += d
            seg = work / f"seg{i}.mp4"
            run(["ffmpeg", "-y", "-v", "error",
                 "-loop", "1", "-framerate", "30", "-i", str(frame),
                 "-i", str(au),
                 "-filter_complex", f"[1:a]apad=pad_dur={a.tail},aresample=48000[a]",
                 "-map", "0:v", "-map", "[a]",
                 "-c:v", "libx264", "-preset", "medium", "-crf", "20",
                 "-pix_fmt", "yuv420p", "-r", "30",
                 "-c:a", "aac", "-b:a", "192k", "-ac", "2",
                 "-t", f"{d:.3f}", str(seg)])
            segments.append(seg)
            print(f"  slide {i}: {au.name}  {d:.1f}s")

        lst = work / "concat.txt"
        lst.write_text("".join(f"file '{s}'\n" for s in segments))
        run(["ffmpeg", "-y", "-v", "error", "-f", "concat", "-safe", "0",
             "-i", str(lst), "-c", "copy", str(a.out)])

        got = duration(a.out)
        print(f"\n完成：{a.out}  {got:.1f} 秒  {Path(a.out).stat().st_size/1e6:.1f} MB")
        if got > 180:
            print(f"⚠ 超過 3 分鐘上限 {got-180:.1f} 秒——請縮短旁白或提高 TTS 語速後重跑")
        else:
            print(f"✓ 在 3 分鐘上限內，剩餘 {180-got:.1f} 秒")
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
