"""用官方 pptx 模板產生 APC_TDC 的 3 分鐘簡報。

模板：VisionChallenge_template-teams.pptx（4 張：Team / Method / Key design decisions / What did not work）
只填內容，不動 master、layout、配色與標題文字。
用法：python make_slides.py <template.pptx> <out.pptx> <figure.png>
"""
import sys, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.oxml.xmlchemy import OxmlElement

TEMPLATE, OUT, FIG = sys.argv[1], sys.argv[2], sys.argv[3]
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0xB4, 0x54, 0x1E)
MUTED = RGBColor(0x55, 0x55, 0x55)


def body_of(slide):
    """回傳內容 placeholder（標題以外那一個）。"""
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith("Inhaltsplatzhalter"):
            return sh
    return None


def place(shape, left, top, width, height):
    """模板把內容框壓在左半邊（9.71"），但版面 1.4" 以下整片是空的（layout 本身就給 12.95"）。"""
    shape.left, shape.top, shape.width, shape.height = (
        Inches(left), Inches(top), Inches(width), Inches(height))


def no_bullet(paragraph):
    """母片的內容樣式帶自動編號，會在每段前面長出「1.」。塞 <a:buNone/> 關掉。"""
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("buChar", "buAutoNum", "buNone"):
        for el in pPr.findall(qn("a:" + tag)):
            pPr.remove(el)
    pPr.append(OxmlElement("a:buNone"))


def fill(shape, blocks, size=16, gap=6):
    """blocks: [(粗體前導, 其餘文字)]；前導為 None 時整行同色。"""
    tf = shape.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE      # 不讓 PowerPoint 自行縮字
    except Exception:
        pass
    tf.clear()
    for i, (lead, rest) in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        no_bullet(p)
        if lead:
            r = p.add_run(); r.text = lead
            r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = ACCENT
        if rest:
            r = p.add_run(); r.text = rest
            r.font.bold = False; r.font.size = Pt(size); r.font.color.rgb = INK


prs = Presentation(TEMPLATE)
SW = prs.slide_width / 914400.0
s1, s2, s3, s4 = list(prs.slides)[:4]

# ---------------------------------------------------------------- slide 1 Team
place(body_of(s1), 0.93, 2.0, 17.1, 4.2)
fill(body_of(s1), [
    ("APC_TDC", "  ·  Tze-Hsiang Tang  ·  Schneider Electric Taiwan Co., Ltd."),
    (None, "A one-person team. We entered both categories."),
    ("", ""),
    ("Category 1", "  weakly supervised tool detection      mAP  0.4930"),
    ("Category 2", "  surgical video question answering     BERTScore-F1  0.6191"),
], size=24, gap=14)

for sh in s1.shapes:                       # 換掉模板的假引用
    if sh.has_text_frame and "fake source" in sh.text_frame.text:
        tf = sh.text_frame; tf.clear()
        no_bullet(tf.paragraphs[0])
        r = tf.paragraphs[0].add_run()
        r.text = "Code, weights and report:  github.com/seantangth/miccai2026-surgvu-surgical-tool-detection-vqa"
        r.font.size = Pt(14); r.font.color.rgb = MUTED

# -------------------------------------------------------------- slide 2 Method
place(body_of(s2), 0.93, 1.9, 17.1, 1.1)
fill(body_of(s2), [
    ("The labels we are given are tool presence, not boxes.",
     "  We turn them into a gate on a teacher ensemble, and let the accepted frames train the next teacher."),
], size=20)

img_w = Inches(15.2)
s2.shapes.add_picture(FIG, Inches((SW - 15.2) / 2), Inches(2.95), width=img_w)

cap = s2.shapes.add_textbox(Inches((SW - 15.2) / 2), Inches(9.62), img_w, Inches(0.9))
fill(cap, [
    ("Accept a frame only if", " the per-class box count equals the installed count exactly and every box has conf ≥ 0.35."
                               "   26% of frames pass; retraining the teacher on them raises that to 41%."),
], size=16, gap=0)

# --------------------------------------------- slide 3 Key design decisions
place(body_of(s3), 0.93, 1.9, 17.1, 8.4)
fill(body_of(s3), [
    ("1.  Scale beats recipe.",
     "  Growing the pseudo-label source from 20 to 280 videos: +0.052. Every recipe change we tried: at most 0.006."),
    ("2.  Let the weak label do the filtering.",
     "  A count-exact match against tools.csv is model-independent, so iterative pseudo-labelling converges instead of drifting."),
    ("3.  Never size-filter pseudo-labels.",
     "  Clipping box areas to P2–P98 removes the near- and far-field samples that carry the scale diversity, and AP75 drops."),
    ("4.  Never keep a frame whose visible tool has no box.",
     "  Three classes were excluded from the targets while their images stayed: 65,182 frames taught the detector"
     " that a visible prograsp is background, against 711 positive boxes. Deleting those frames: 0.000 → 0.255 detection."),
    ("5.  The remaining headroom is AP75, not recall.",
     "  Needle driver sits at AP50 0.93 but AP75 0.51. Training at 800 px: +0.0275 AP75. Two-model fusion: +0.0376 AP75,"
     " measured on a pair of 640 px models. Our submission fuses 640 with 800, a combination we could not measure locally."),
], size=20, gap=17)

# ------------------------------------------- slide 4 What did not work
place(body_of(s4), 0.93, 1.9, 17.1, 8.4)
fill(body_of(s4), [
    ("Inference-time search.",
     "  39 variants of TTA, confidence, top-k, multi-scale and NMS. Ceiling: +0.005."),
    ("Temporal post-processing.",
     "  Ten variants of score smoothing, box smoothing and gap filling. All negative, −0.004 to −0.06."),
    ("Training on the official ground truth alone.",
     "  0.4024 against 0.5905 for mixed training. Fine-tuning on it afterwards also hurt: needle-driver AP75 0.51 → 0.36."
     "  Five videos is a narrower scale prior than the pseudo-labels provide."),
    ("Reading the robot interface with OCR to relabel.",
     "  It fixed the four never-detected classes on a proxy metric, but cost −0.0127 mAP on the common classes."),
    ("Trusting a proxy metric.",
     "  That same relabelling moved tool-presence set-F1 from 0.578 to 0.819 while moving real mAP the other way."
     "  Two rulers, opposite directions."),
    ("", ""),
    ("And the honest caveat:",
     "  our ablations hold out the student, not the teacher, so read the paired differences and not the absolute values."),
], size=19, gap=15)


# ---------------------------------- 備忘稿：直接讀 narration/，與 TTS 稿保持單一來源
NARR = pathlib.Path(__file__).resolve().parent / "narration"
if NARR.is_dir():
    total = 0
    for i, sl in enumerate(prs.slides, 1):
        f = NARR / f"slide{i}.txt"
        if not f.exists():
            continue
        text = " ".join(f.read_text().split())
        total += len(text.split())
        sl.notes_slide.notes_text_frame.text = text
    print(f"備忘稿取自 narration/：合計 {total} 字 ≈ {total/2.5:.0f}–{total/2.7:.0f} 秒")
else:
    print("找不到 narration/，備忘稿留空")

prs.save(OUT)
print("wrote", OUT)
